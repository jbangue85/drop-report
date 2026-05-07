import openpyxl
from typing import List, Dict, Any

# Exact column names from Dropi XLSX → SQLite field names
COLUMN_MAP: Dict[str, str] = {
    "FECHA DE REPORTE":              "fecha_reporte",
    "ID":                            "id",
    "HORA":                          "hora",
    "FECHA":                         "fecha",
    "NOMBRE CLIENTE":                "nombre_cliente",
    "TELÉFONO":                      "telefono",
    "EMAIL":                         "email",
    "TIPO DE IDENTIFICACION":        "tipo_identificacion",
    "NRO DE IDENTIFICACION":         "nro_identificacion",
    "NÚMERO GUIA":                   "numero_guia",
    "ESTATUS":                       "estatus",
    "TIPO DE ENVIO":                 "tipo_envio",
    "DEPARTAMENTO DESTINO":          "departamento_destino",
    "CIUDAD DESTINO":                "ciudad_destino",
    "DIRECCION":                     "direccion",
    "NOTAS":                         "notas",
    "TRANSPORTADORA":                "transportadora",
    "TOTAL DE LA ORDEN":             "total_orden",
    "GANANCIA":                      "ganancia",
    "PRECIO FLETE":                  "precio_flete",
    "COSTO DEVOLUCION FLETE":        "costo_devolucion_flete",
    "COMISION":                      "comision",
    "% COMISION DE LA PLATAFORMMA":  "pct_comision",
    "PRECIO PROVEEDOR":              "precio_proveedor",
    "PRECIO PROVEEDOR X CANTIDAD":   "precio_proveedor_x_cantidad",
    "PRODUCTO ID":                   "producto_id",
    "SKU":                           "sku",
    "VARIACION ID":                  "variacion_id",
    "PRODUCTO":                      "producto",
    "VARIACION":                     "variacion",
    "CANTIDAD":                      "cantidad",
    "NOVEDAD":                       "novedad",
    "FUE SOLUCIONADA LA NOVEDAD":    "novedad_solucionada",
    "HORA DE NOVEDAD":               "hora_novedad",
    "FECHA DE NOVEDAD":              "fecha_novedad",
    "SOLUCIÓN":                      "solucion",
    "HORA DE SOLUCIÓN":              "hora_solucion",
    "FECHA DE SOLUCIÓN":             "fecha_solucion",
    "OBSERVACIÓN":                   "observacion",
    "HORA DE ÚLTIMO MOVIMIENTO":     "hora_ultimo_movimiento",
    "FECHA DE ÚLTIMO MOVIMIENTO":    "fecha_ultimo_movimiento",
    "ÚLTIMO MOVIMIENTO":             "ultimo_movimiento",
    "CONCEPTO ÚLTIMO MOVIMIENTO":    "concepto_ultimo_movimiento",
    "UBICACIÓN DE ÚLTIMO MOVIMIENTO":"ubicacion_ultimo_movimiento",
    "VENDEDOR":                      "vendedor",
    "TIPO DE TIENDA":                "tipo_tienda",
    "TIENDA":                        "tienda",
    "ID DE ORDEN DE TIENDA":         "id_orden_tienda",
    "NUMERO DE PEDIDO DE TIENDA":    "numero_pedido_tienda",
    "TAGS":                          "tags",
    "FECHA GUIA GENERADA":           "fecha_guia_generada",
    "CONTADOR DE INDEMNIZACIONES":   "contador_indemnizaciones",
    "CONCEPTO ÚLTIMA INDENMIZACIÓN": "concepto_ultima_indemnizacion",
}

REQUIRED_DROPI_PRODUCT_HEADERS = {
    "ID",
    "FECHA",
    "NOMBRE CLIENTE",
    "TELÉFONO",
    "NÚMERO GUIA",
    "ESTATUS",
    "DEPARTAMENTO DESTINO",
    "CIUDAD DESTINO",
    "DIRECCION",
    "TRANSPORTADORA",
    "TOTAL DE LA ORDEN",
    "PRODUCTO ID",
    "PRODUCTO",
    "CANTIDAD",
    "PRECIO PROVEEDOR X CANTIDAD",
}


class InvalidDropiFileError(ValueError):
    pass


class InvalidMetaFileError(ValueError):
    pass


def _normalize_header(value) -> str:
    return str(value or "").strip()


def _validate_dropi_product_headers(headers) -> None:
    present = {_normalize_header(h) for h in headers if _normalize_header(h)}
    missing = sorted(REQUIRED_DROPI_PRODUCT_HEADERS - present)
    if missing:
        raise InvalidDropiFileError(
            "El archivo Dropi no tiene la estructura esperada de ordenes_productos. "
            "Faltan columnas obligatorias: " + ", ".join(missing)
        )


def _validate_meta_headers(headers) -> str:
    present = {_normalize_header(h) for h in headers if _normalize_header(h)}
    spend_cols = sorted(h for h in present if "Importe gastado" in h)
    required = {
        "Inicio del informe",
        "Fin del informe",
        "Nombre de la campaña",
        "Resultados",
        "Indicador de resultado",
    }
    missing = sorted(required - present)
    if not spend_cols:
        missing.append("Importe gastado")
    if missing:
        raise InvalidMetaFileError(
            "El archivo Meta Ads no tiene la estructura esperada del reporte de campañas. "
            "Faltan columnas obligatorias: " + ", ".join(missing)
        )
    return spend_cols[0]


def _normalize_date(value) -> "Optional[str]":
    """Convert DD-MM-YYYY to YYYY-MM-DD for SQLite date ordering."""
    if not value:
        return None
    s = str(value).strip()
    if len(s) == 10 and s[2] == "-":
        # DD-MM-YYYY → YYYY-MM-DD
        d, m, y = s.split("-")
        return f"{y}-{m}-{d}"
    return s


def parse_xlsx(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parse Dropi XLSX bytes → list of dicts ready for SQLite upsert."""
    import io
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    raw_headers = rows[0]
    _validate_dropi_product_headers(raw_headers)

    # Build index: db_field → column_index (skip unknown columns)
    col_index: Dict[str, int] = {}
    for i, h in enumerate(raw_headers):
        header = _normalize_header(h)
        if header in COLUMN_MAP:
            db_field = COLUMN_MAP[header]
            col_index[db_field] = i

    date_fields = {
        "fecha_reporte", "fecha", "fecha_novedad",
        "fecha_solucion", "fecha_ultimo_movimiento", "fecha_guia_generada",
    }

    records = []
    for row in rows[1:]:
        if not any(row):
            continue
        rec: Dict[str, Any] = {"source_file": filename}
        for field, idx in col_index.items():
            val = row[idx] if idx < len(row) else None
            if field in date_fields:
                val = _normalize_date(val)
            rec[field] = val
        # Must have an id to upsert
        if rec.get("id") is not None:
            records.append(rec)

    return records


def upsert_records(conn, records: List[Dict[str, Any]]) -> int:
    """INSERT OR REPLACE all records. Returns count inserted/updated."""
    if not records:
        return 0

    all_fields = list({k for r in records for k in r.keys()})
    placeholders = ", ".join(["?" for _ in all_fields])
    cols = ", ".join(all_fields)
    sql = f"INSERT OR REPLACE INTO orders ({cols}) VALUES ({placeholders})"

    data = []
    for rec in records:
        data.append(tuple(rec.get(f) for f in all_fields))

    conn.executemany(sql, data)
    conn.commit()
    return len(records)


def parse_meta_csv(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Parse Meta Ads CSV bytes → list of dicts for meta_ads_spend table."""
    import csv
    import io
    
    text = file_bytes.decode('utf-8-sig', errors='replace')
    reader = csv.DictReader(io.StringIO(text))
    spend_col = _validate_meta_headers(reader.fieldnames or [])
    
    records = []
    for row in reader:
        fecha_raw = row.get("Inicio del informe", "").strip()
        # Handle various Meta date formats (YYYY-MM-DD, DD/MM/YYYY, etc.)
        fecha = fecha_raw
        if "/" in fecha_raw:
            parts = fecha_raw.split("/")
            if len(parts) == 3:
                # If first part is 4 digits, it's likely YYYY/MM/DD
                if len(parts[0]) == 4:
                    fecha = f"{parts[0]}-{parts[1]}-{parts[2]}"
                # Else assume DD/MM/YYYY
                else:
                    fecha = f"{parts[2]}-{parts[1]}-{parts[0]}"
        elif len(fecha_raw) == 10 and fecha_raw[2] == "-":
            # DD-MM-YYYY -> YYYY-MM-DD
            d, m, y = fecha_raw.split("-")
            fecha = f"{y}-{m}-{d}"
            
        campaign = row.get("Nombre de la campaña", "").strip()
        spend_str = row[spend_col].strip() if row.get(spend_col) else "0"
        results_str = row.get("Resultados", "0").strip()
        
        if not fecha or not campaign:
            continue
            
        try:
            spend = float(spend_str.replace(',', '.')) if spend_str else 0.0
        except ValueError:
            spend = 0.0
            
        try:
            results = int(results_str) if results_str else 0
        except ValueError:
            results = 0
            
        if spend > 0 or results > 0:
            records.append({
                "fecha": fecha,
                "campaign_name": campaign,
                "spend": spend,
                "results": results
            })
            
    return records


def upsert_meta_spend(conn, records: List[Dict[str, Any]]) -> int:
    """INSERT OR REPLACE meta ads spend and auto-map new campaigns."""
    if not records:
        return 0

    # 1. Insert spend
    sql = """
        INSERT OR REPLACE INTO meta_ads_spend 
        (fecha, campaign_name, spend, results) 
        VALUES (?, ?, ?, ?)
    """
    data = [(r["fecha"], r["campaign_name"], r["spend"], r["results"]) for r in records]
    conn.executemany(sql, data)

    # 2. Find unmapped campaigns
    all_campaigns = list(set(r["campaign_name"] for r in records))
    placeholders = ",".join("?" * len(all_campaigns))
    mapped = conn.execute(f"SELECT campaign_name FROM campaign_map WHERE campaign_name IN ({placeholders})", all_campaigns).fetchall()
    mapped_set = set(r["campaign_name"] for r in mapped)
    
    unmapped = [c for c in all_campaigns if c not in mapped_set]
    if unmapped:
        # Get all products
        products = [r[0] for r in conn.execute("SELECT DISTINCT producto FROM orders WHERE producto IS NOT NULL").fetchall()]
        
        import re
        import unicodedata
        def normalize(s):
            if not s: return ""
            s = unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode("utf-8")
            s = s.lower()
            return re.sub(r'[^a-z0-9]', ' ', s)
            
        map_data = []
        for c in unmapped:
            nc = set(normalize(c).split())
            best_p = None
            best_score = 0
            for p in products:
                np = set(normalize(p).split())
                score = len(nc.intersection(np))
                if score > best_score and score > 0:
                    best_score = score
                    best_p = p
            
            # Insert even if best_p is None, so it shows up in UI to be mapped
            map_data.append((c, best_p))
            
        conn.executemany("INSERT OR REPLACE INTO campaign_map (campaign_name, producto) VALUES (?, ?)", map_data)

    conn.commit()
    return len(records)
